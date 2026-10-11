"""
Tạo file Excel mẫu danh sách in áo: apps/web/public/tai-ve/mau-danh-sach-in-ao-yala.xlsx
Chạy lại khi đổi cột/size:  python scripts/gen/roster-template.py   (cần: pip install openpyxl)

Quy tắc (khớp bộ đọc roster ở packages/shared/src/roster.ts + apps/web/src/lib/xlsx.ts):
- Sheet ĐẦU TIÊN là "Danh sách", dòng 1 là tiêu đề: Tên in | Số áo | Size | Ghi chú.
- Sheet đầu không có dòng ví dụ, không có công thức -> tải lên là ra đúng số áo, không dòng thừa.
- Ví dụ, hướng dẫn, tổng size nằm ở các sheet sau (bộ đọc chỉ lấy sheet đầu).
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parents[2] / "apps/web/public/tai-ve/mau-danh-sach-in-ao-yala.xlsx"
SIZES = ["S", "M", "L", "XL", "2XL", "3XL"]
ROWS = 200
HEAD = ["Tên in", "Số áo", "Size", "Ghi chú"]
WIDTH = [28, 10, 10, 34]

ORANGE = "E4570B"
INK = "1D1D1F"
LINE = Side(style="thin", color="D9D6D0")
head_font = Font(bold=True, color="FFFFFF", size=12)
head_fill = PatternFill("solid", fgColor=INK)
border = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)


def table_sheet(ws, rows):
    ws.append(HEAD)
    for c in ws[1]:
        c.font, c.fill, c.alignment, c.border = head_font, head_fill, Alignment(horizontal="center", vertical="center"), border
    ws.row_dimensions[1].height = 24
    for i, w in enumerate(WIDTH):
        ws.column_dimensions["ABCD"[i]].width = w
    ws.freeze_panes = "A2"
    for r in rows:
        ws.append(r)
    # Số áo dạng chữ để giữ số 0 đầu (07, 00)
    for r in range(2, ROWS + 2):
        ws.cell(r, 2).number_format = "@"
        for col in range(1, 5):
            ws.cell(r, col).border = border
        ws.cell(r, 2).alignment = ws.cell(r, 3).alignment = Alignment(horizontal="center")


wb = Workbook()

# 1) Danh sách (sheet đầu – YALA đọc sheet này)
ws = wb.active
ws.title = "Danh sách"
table_sheet(ws, [])
dv_size = DataValidation(type="list", formula1=f'"{",".join(SIZES)}"', allow_blank=True, errorStyle="warning",
                         errorTitle="Size", error="Size không có trong danh sách S–3XL. Vẫn giữ nếu bạn đã hỏi xưởng.",
                         promptTitle="Size", prompt="Chọn S, M, L, XL, 2XL, 3XL", showErrorMessage=True, showInputMessage=True)
dv_num = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="3", allow_blank=True, errorStyle="warning",
                        errorTitle="Số áo", error="Số áo tối đa 3 chữ số (0–999).", showErrorMessage=True)
dv_name = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="30", allow_blank=True, errorStyle="warning",
                         errorTitle="Tên in", error="Tên in tối đa 30 ký tự.", showErrorMessage=True)
for dv, col in ((dv_size, "C"), (dv_num, "B"), (dv_name, "A")):
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{ROWS + 1}")

# 2) Ví dụ
ex = wb.create_sheet("Ví dụ")
table_sheet(ex, [
    ["NGUYỄN AN", "10", "M", ""],
    ["TRẦN BÌNH", "07", "L", ""],
    ["LÊ CHI", "", "S", "Không in số"],
    ["PHẠM DŨNG", "99", "XL", "Đội trưởng"],
    ["", "1", "2XL", "Chỉ in số, không in tên"],
])

# 3) Tổng size – tự đếm từ sheet Danh sách
sm = wb.create_sheet("Tổng size")
sm.append(["Size", "Số áo"])
for c in sm[1]:
    c.font, c.fill, c.border = head_font, head_fill, border
for s in SIZES:
    sm.append([s, f"=COUNTIF('Danh sách'!C2:C{ROWS + 1},\"{s}\")"])
sm.append(["Tổng", f"=COUNTA('Danh sách'!C2:C{ROWS + 1})"])
last = sm.max_row
for r in range(2, last + 1):
    for col in (1, 2):
        sm.cell(r, col).border = border
sm.cell(last, 1).font = sm.cell(last, 2).font = Font(bold=True, color=ORANGE)
sm.column_dimensions["A"].width = 12
sm.column_dimensions["B"].width = 12

# 4) Hướng dẫn
hd = wb.create_sheet("Hướng dẫn")
hd.column_dimensions["A"].width = 100
lines = [
    ("MẪU DANH SÁCH IN ÁO – YALA (yala.vn)", Font(bold=True, size=14, color=ORANGE)),
    ("", None),
    ("1. Điền ở sheet “Danh sách”: mỗi dòng là 1 áo. Không đổi tên cột ở dòng 1.", None),
    ("2. Tên in: tên in trên lưng áo, tối đa 30 ký tự. Để trống nếu không in tên.", None),
    ("3. Số áo: 0–999. Muốn giữ số 0 đầu (07) thì cứ gõ 07 – cột đã để dạng chữ. Để trống nếu không in số.", None),
    ("4. Size: chọn trong ô thả xuống (S, M, L, XL, 2XL, 3XL). Xem cách đo tại yala.vn/bang-size-ao", None),
    ("5. Ghi chú: yêu cầu riêng cho áo đó (đội trưởng, không in số…).", None),
    ("6. Sheet “Tổng size” tự đếm số áo theo từng size để bạn kiểm tra trước khi đặt.", None),
    ("", None),
    ("Tải lên: vào trang sản phẩm áo trên yala.vn → tab “Đặt cho nhóm” → “Tải file Excel / CSV”.", Font(bold=True)),
    ("Hoặc bôi đen các cột trong Excel / Google Sheets → Ctrl+C → “Dán từ Excel”.", None),
    ("Dùng Google Sheets: Tệp → Nhập → Tải lên file này; xong Tải xuống → Microsoft Excel (.xlsx).", None),
    ("", None),
    ("Đơn nhóm từ 10 sản phẩm. Ít hơn vẫn đặt được, tính như mua lẻ.", None),
    ("Cần hỗ trợ: Zalo/Hotline 0971 808 330 · contact@yala.vn", None),
]
for text, font in lines:
    hd.append([text])
    if font:
        hd.cell(hd.max_row, 1).font = font
    hd.cell(hd.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")

wb.active = 0
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"OK {OUT} ({OUT.stat().st_size} bytes)")
